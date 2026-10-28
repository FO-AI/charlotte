from datetime import date
from decimal import Decimal, InvalidOperation, ROUND_HALF_UP
from io import BytesIO
import re
from typing import Any, Dict, List, Mapping, Optional, Sequence, Tuple

from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

# Banking asked for Provider text length between 1 and 30 characters.
PROVIDER_MIN_LENGTH = 1
PROVIDER_MAX_LENGTH = 30
# Academic year label for Fall 2026 – Summer 2 2027 (not calendar year of the fall term).
DEFAULT_AID_YEAR = "2027"
ALLOWED_AID_TERMS = frozenset({"F", "S", "F/S", "SS1", "SS2"})
DEFAULT_AID_TERM = "F"
# UNC student PID: exactly nine digits. Approval / check numbers must not become Excel rows.
_PID_DIGIT_COUNT = 9


class OutsideScholarshipsDataLoader:
    """Build an Excel workbook from outside-scholarship extraction or review output."""

    HEADERS = [
        "PID",
        "Amount",
        "Check #",
        "Name",
        "Active Directory Name",
        "Aid year",
        "Aid term",
        "Provider",
        "Scholarship name",
    ]

    REVIEW_HEADERS = HEADERS + [
        "Edited",
        "Verified",
        "Extracted values",
    ]

    def __init__(
        self,
        aid_year: Optional[str] = None,
        aid_term: Optional[str] = None,
        active_directory_names: Optional[Mapping[str, str]] = None,
        reviewed_by: Optional[str] = None,
    ):
        self.aid_year = self._normalize_aid_year(aid_year)
        self.aid_term = self._normalize_aid_term(aid_term)
        self.active_directory_names = active_directory_names or {}
        self.reviewed_by = (reviewed_by or "").strip()
        self.report_date = date.today()

    @staticmethod
    def _normalize_aid_year(aid_year: Optional[str]) -> str:
        value = str(aid_year).strip() if aid_year is not None else ""
        if not value:
            return DEFAULT_AID_YEAR
        if not value.isdigit() or len(value) != 4:
            raise ValueError("Aid year must be a 4-digit year (for example: 2027).")
        return value

    @staticmethod
    def _normalize_aid_term(aid_term: Optional[str]) -> str:
        value = (aid_term or "").strip().upper() or DEFAULT_AID_TERM
        if value not in ALLOWED_AID_TERMS:
            allowed = ", ".join(sorted(ALLOWED_AID_TERMS))
            raise ValueError(f"Aid term must be one of: {allowed}.")
        return value

    def _row_aid_term(self, reviewed: Mapping[str, Any]) -> str:
        """Use the term read for this check. Fall back to the upload-screen term when it has none."""
        raw = self._text(reviewed.get("aid_term"))
        if raw:
            normalized = raw.strip().upper()
            if normalized in ALLOWED_AID_TERMS:
                return normalized
        return self.aid_term

    @staticmethod
    def _to_decimal_amount(value: Any) -> Decimal:
        if value is None:
            return Decimal("0")

        text = str(value).strip()
        if not text:
            return Decimal("0")

        # Match the review UI: strip currency, thousands separators, and internal spaces
        # so values like "1 250.00" export as 1250.00 rather than quietly becoming $0.
        cleaned = text.replace("$", "").replace(",", "").replace(" ", "")
        try:
            return Decimal(cleaned)
        except InvalidOperation:
            return Decimal("0")

    @staticmethod
    def _parse_well_formed_amount(value: Any) -> Optional[Decimal]:
        if value is None:
            return None

        text = str(value).strip()
        if not text:
            return None

        cleaned = text.replace("$", "").replace(",", "").replace(" ", "")
        if not re.fullmatch(r"\d+(?:\.\d{1,2})?", cleaned):
            return None
        try:
            return Decimal(cleaned)
        except InvalidOperation:
            return None

    @staticmethod
    def _text(value: Any) -> str:
        if value is None:
            return ""
        return str(value).strip()

    @staticmethod
    def _excel_safe_text(value: Any) -> str:
        """Neutralize formula injection when a cell value starts with =, +, -, or @."""
        text = OutsideScholarshipsDataLoader._text(value)
        if text and text[0] in {"=", "+", "-", "@"}:
            return f"'{text}"
        return text

    @staticmethod
    def _normalize_pid_digits(value: Any) -> str:
        return "".join(ch for ch in str(value or "") if ch.isdigit())

    @classmethod
    def _nine_digit_pid(cls, value: Any) -> Optional[str]:
        """Return the PID's digits when it is exactly nine digits; otherwise None."""
        digits = cls._normalize_pid_digits(value)
        return digits if len(digits) == _PID_DIGIT_COUNT else None

    @classmethod
    def _nine_digit_pids_from_value(cls, value: Any) -> List[str]:
        """Nine-digit PIDs in one cell/field, including when an approval number shares the string."""
        text = cls._text(value)
        if not text:
            return []
        whole = cls._nine_digit_pid(text)
        if whole:
            return [whole]

        found: List[str] = []
        seen = set()
        for token in re.split(r"[\s,;/|]+", text):
            pid = cls._nine_digit_pid(token)
            if not pid or pid in seen:
                continue
            seen.add(pid)
            found.append(pid)
        if found:
            return found

        for match in re.finditer(rf"(?<!\d)\d{{{_PID_DIGIT_COUNT}}}(?!\d)", text):
            pid = match.group(0)
            if pid not in seen:
                seen.add(pid)
                found.append(pid)
        return found

    @staticmethod
    def _pid_entries(check: Dict[str, Any], key: str) -> List[Dict[str, Any]]:
        raw = check.get(key)
        if not isinstance(raw, list):
            return []
        entries = []
        for item in raw:
            if isinstance(item, dict):
                entries.append(item)
            elif item is not None and str(item).strip():
                entries.append({"pid": str(item).strip()})
        return entries

    @classmethod
    def _nine_digit_pid_entries(cls, entries: Sequence[Dict[str, Any]]) -> List[Dict[str, Any]]:
        """Keep one entry per distinct nine-digit PID; drop approval numbers and other junk."""
        filtered: List[Dict[str, Any]] = []
        seen = set()
        for entry in entries:
            for pid in cls._nine_digit_pids_from_value(entry.get("pid")):
                if pid in seen:
                    continue
                seen.add(pid)
                filtered.append({**entry, "pid": pid})
        return filtered

    @classmethod
    def _audit_diff(
        cls,
        extracted: Dict[str, Any],
        reviewed: Dict[str, Any],
        extracted_pid_entries: Sequence[Dict[str, Any]],
        reviewed_pid_entries: Sequence[Dict[str, Any]],
    ) -> Tuple[bool, str]:
        changes: List[str] = []
        field_labels = (
            ("amount", "Amount"),
            ("name", "Name"),
            ("check_number", "Check number"),
            ("provider", "Provider"),
            ("scholarship_name", "Scholarship name"),
            ("aid_term", "Aid term"),
        )
        for field, label in field_labels:
            extracted_value = cls._text(extracted.get(field))
            reviewed_value = cls._text(reviewed.get(field))
            if extracted_value != reviewed_value:
                changes.append(f"{label}: {extracted_value or '(blank)'}")

        extracted_set = [entry["pid"] for entry in extracted_pid_entries if entry.get("pid")]
        reviewed_set = [entry["pid"] for entry in reviewed_pid_entries if entry.get("pid")]
        for pid in reviewed_set:
            if pid not in extracted_set:
                changes.append(f"Added PID: {pid}")
        for pid in extracted_set:
            if pid not in reviewed_set:
                changes.append(f"Removed PID: {pid}")

        extracted_amounts = {
            entry["pid"]: cls._text(entry.get("amount"))
            for entry in extracted_pid_entries
            if entry.get("pid")
        }
        reviewed_amounts = {
            entry["pid"]: cls._text(entry.get("amount"))
            for entry in reviewed_pid_entries
            if entry.get("pid")
        }
        for pid in reviewed_set:
            if extracted_amounts.get(pid, "") != reviewed_amounts.get(pid, ""):
                changes.append(f"PID {pid} amount: {extracted_amounts.get(pid) or '(blank)'}")

        edited = bool(changes)
        return edited, "; ".join(changes)

    @classmethod
    def _row_amounts_for_reviewed_pids(
        cls,
        reviewed_pid_entries: Sequence[Dict[str, Any]],
        check_amount: Decimal,
    ) -> Optional[Dict[str, Decimal]]:
        non_blank_entries = [entry for entry in reviewed_pid_entries if entry.get("pid")]
        if len(non_blank_entries) <= 1:
            return {}

        amounts: Dict[str, Decimal] = {}
        for entry in non_blank_entries:
            pid = entry.get("pid")
            amount = cls._parse_well_formed_amount(entry.get("amount"))
            if pid is None or amount is None:
                return None
            amounts[pid] = amount

        split_total = sum(amounts.values(), Decimal("0"))
        if abs(split_total - check_amount) > Decimal("0.01"):
            return None
        return amounts

    def _expand_reviewed_rows(self, checks: List[Dict[str, Any]]) -> List[List[Any]]:
        rows: List[List[Any]] = []

        for check in checks:
            extracted = check.get("extracted") if isinstance(check.get("extracted"), dict) else {}
            reviewed = check.get("reviewed") if isinstance(check.get("reviewed"), dict) else check
            verified = bool(check.get("verified"))

            amount_decimal = self._to_decimal_amount(reviewed.get("amount"))
            check_number = self._excel_safe_text(reviewed.get("check_number") or "")
            name = self._excel_safe_text(reviewed.get("name") or "")
            extracted_name = self._excel_safe_text(extracted.get("name") or "")
            provider = self._excel_safe_text(reviewed.get("provider") or "")
            scholarship_name = self._excel_safe_text(reviewed.get("scholarship_name") or "")
            extracted_scholarship_name = self._excel_safe_text(extracted.get("scholarship_name") or "")
            check_level_name_edited = name != extracted_name
            check_level_scholarship_edited = scholarship_name != extracted_scholarship_name

            reviewed_pid_entries = self._pid_entries(reviewed, "pids")
            if not reviewed_pid_entries:
                # Fall back to pid_list shape if present.
                raw_list = reviewed.get("pid_list")
                if isinstance(raw_list, list) and raw_list:
                    reviewed_pid_entries = [{"pid": pid} for pid in raw_list]

            extracted_pid_entries = self._pid_entries(extracted, "pids")
            if not extracted_pid_entries:
                raw_list = extracted.get("pid_list")
                if isinstance(raw_list, list):
                    extracted_pid_entries = [{"pid": pid} for pid in raw_list]

            # Issue #13: strip approval numbers from extracted PIDs only. Reviewed values are
            # validated in build_reviewed_excel_bytes and must not be silently dropped.
            extracted_pid_entries = self._nine_digit_pid_entries(extracted_pid_entries)

            normalized_reviewed: List[Dict[str, Any]] = []
            seen_reviewed = set()
            for entry in reviewed_pid_entries:
                text = self._text(entry.get("pid"))
                if not text:
                    continue
                digits = self._normalize_pid_digits(text) or text
                if digits in seen_reviewed:
                    continue
                seen_reviewed.add(digits)
                normalized_reviewed.append({**entry, "pid": digits})
            reviewed_pid_entries = normalized_reviewed or [{"pid": ""}]

            extracted_pids = [entry["pid"] for entry in extracted_pid_entries]

            split_amounts = self._row_amounts_for_reviewed_pids(reviewed_pid_entries, amount_decimal)

            edited, extracted_values = self._audit_diff(
                extracted,
                reviewed,
                extracted_pid_entries,
                reviewed_pid_entries,
            )
            extracted_values = self._excel_safe_text(extracted_values)

            for entry in reviewed_pid_entries:
                pid_text = entry.get("pid") or ""
                row_amount_decimal = (
                    amount_decimal if split_amounts is None else split_amounts.get(pid_text, amount_decimal)
                )
                row_amount_value = float(row_amount_decimal.quantize(Decimal("0.01"), rounding=ROUND_HALF_UP))
                if check_level_name_edited:
                    row_name = name
                else:
                    row_name = self._excel_safe_text(entry.get("name") or name)
                if check_level_scholarship_edited:
                    row_scholarship_name = scholarship_name
                else:
                    row_scholarship_name = self._excel_safe_text(entry.get("scholarship_name") or scholarship_name)
                # Server-resolved AD names only; never trust client-sent active_directory.
                ad_name = self.active_directory_names.get(pid_text, "") if pid_text else ""
                row_edited = edited
                # An added PID alone marks the row edited (already in extracted_values).
                if pid_text and pid_text not in extracted_pids:
                    row_edited = True
                rows.append(
                    [
                        self._excel_safe_text(pid_text),
                        row_amount_value,
                        check_number,
                        row_name,
                        self._excel_safe_text(ad_name),
                        self.aid_year,
                        self._row_aid_term(reviewed),
                        provider,
                        row_scholarship_name,
                        "Yes" if row_edited else "",
                        "Yes" if verified else "",
                        extracted_values if row_edited else "",
                    ]
                )

        return rows

    @classmethod
    def _validate_reviewed_pids(cls, checks: Sequence[Dict[str, Any]]) -> None:
        """Reject non-blank reviewed PIDs that are not exactly nine digits."""
        for check in checks:
            reviewed = check.get("reviewed") if isinstance(check.get("reviewed"), dict) else check
            check_number = cls._text(reviewed.get("check_number")) or str(
                check.get("check_index") or "?"
            )
            entries = cls._pid_entries(reviewed, "pids")
            if not entries:
                raw_list = reviewed.get("pid_list")
                if isinstance(raw_list, list) and raw_list:
                    entries = [{"pid": pid} for pid in raw_list]
            for entry in entries:
                text = cls._text(entry.get("pid"))
                if not text:
                    continue
                if cls._nine_digit_pid(text) is None:
                    bad = cls._normalize_pid_digits(text) or text
                    raise ValueError(
                        f"Check {check_number}: PID '{bad}' is not 9 digits. "
                        "Fix or clear the PID before exporting."
                    )

    def build_reviewed_excel_bytes(self, review_payload: Dict[str, Any]) -> BytesIO:
        checks_raw = review_payload.get("checks") if isinstance(review_payload, dict) else []
        checks = [check for check in checks_raw if isinstance(check, dict)] if isinstance(checks_raw, list) else []
        self._validate_reviewed_pids(checks)
        rows = self._expand_reviewed_rows(checks)
        return self._write_workbook(checks, rows, headers=self.REVIEW_HEADERS, include_reviewed_by=True)

    def _write_workbook(
        self,
        checks: List[Dict[str, Any]],
        rows: List[List[Any]],
        headers: Sequence[str],
        include_reviewed_by: bool,
    ) -> BytesIO:
        total_amount = Decimal("0")
        for check in checks:
            reviewed = check.get("reviewed") if isinstance(check.get("reviewed"), dict) else check
            total_amount += self._to_decimal_amount(reviewed.get("amount"))

        workbook = Workbook()
        worksheet = workbook.active
        worksheet.title = "outside_scholarships"

        worksheet["A1"] = "Total amount"
        worksheet["B1"] = float(total_amount.quantize(Decimal("0.01"), rounding=ROUND_HALF_UP))
        worksheet["A2"] = "Check count"
        worksheet["B2"] = len(checks)
        worksheet["A3"] = "Date"
        worksheet["B3"] = self.report_date.isoformat()

        worksheet["A1"].font = Font(bold=True)
        worksheet["A2"].font = Font(bold=True)
        worksheet["A3"].font = Font(bold=True)
        worksheet["B1"].number_format = "$#,##0.00"

        if include_reviewed_by:
            worksheet["A4"] = "Reviewed by"
            worksheet["B4"] = self._excel_safe_text(self.reviewed_by)
            worksheet["A4"].font = Font(bold=True)
            header_row_index = 6
            freeze_pane = "A7"
        else:
            header_row_index = 5
            freeze_pane = "A6"

        for column_index, header in enumerate(headers, start=1):
            cell = worksheet.cell(row=header_row_index, column=column_index, value=header)
            cell.font = Font(bold=True)

        for row_offset, row_values in enumerate(rows, start=1):
            for column_index, value in enumerate(row_values, start=1):
                worksheet.cell(row=header_row_index + row_offset, column=column_index, value=value)

        worksheet.freeze_panes = freeze_pane

        column_widths = {
            "A": 18,
            "B": 12,
            "C": 12,
            "D": 26,
            "E": 26,
            "F": 12,
            "G": 10,
            "H": 26,
            "I": 30,
            "J": 10,
            "K": 10,
            "L": 40,
        }
        for column_name, width in column_widths.items():
            worksheet.column_dimensions[column_name].width = width

        first_data_row = header_row_index + 1
        last_data_row = first_data_row + max(len(rows) - 1, 0)
        for row_index in range(first_data_row, last_data_row + 1):
            worksheet.cell(row=row_index, column=2).number_format = "$#,##0.00"

        try:
            provider_col = list(headers).index("Provider") + 1
        except ValueError:
            provider_col = None
        if provider_col is not None and rows:
            provider_letter = get_column_letter(provider_col)
            # Cover exported rows plus room to paste/extend. allow_blank=True so empty
            # extension rows do not error until text is entered (then 1–30 applies).
            validation_end = max(last_data_row, first_data_row + 499)
            provider_validation = DataValidation(
                type="textLength",
                operator="between",
                formula1=str(PROVIDER_MIN_LENGTH),
                formula2=str(PROVIDER_MAX_LENGTH),
                allow_blank=True,
                showErrorMessage=True,
                errorTitle="Provider length",
                error=(
                    f"Provider must be between {PROVIDER_MIN_LENGTH} and "
                    f"{PROVIDER_MAX_LENGTH} characters."
                ),
                promptTitle="Provider",
                prompt=(
                    f"Enter {PROVIDER_MIN_LENGTH}–{PROVIDER_MAX_LENGTH} characters."
                ),
                showInputMessage=True,
            )
            provider_validation.add(f"{provider_letter}{first_data_row}:{provider_letter}{validation_end}")
            worksheet.add_data_validation(provider_validation)

        output = BytesIO()
        workbook.save(output)
        output.seek(0)
        return output
