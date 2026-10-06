"""Issue #13: approval numbers must not become extra Excel rows beside real PIDs.

A UNC student PID is exactly nine digits. Shorter values (e.g. 1380307) are not PIDs.
"""

from io import BytesIO
from types import SimpleNamespace

import pytest
from openpyxl import load_workbook

from outside_scholarships_fakes import (
    BACK,
    FRONT,
    REVIEW_FIRST_DATA_ROW,
    FakeGraph,
    export_excel,
    preview_payload,
    upload,
)
from services.banking.outside_scholarships.active_directory_names import _normalize_pid as ad_normalize_pid
from services.banking.outside_scholarships.nodes import _dedupe_pids, _normalize_pid
from services.banking.outside_scholarships.service import OutsideScholarshipService
from services.data_loaders.outside_scholarships_json_to_excel import OutsideScholarshipsDataLoader

_PID = "730123456"
_SECOND_PID = "730654321"
_APPROVAL_NUMBER = "1380307"


def test_normalize_pid_keeps_only_nine_digits():
    assert _normalize_pid(_PID) == _PID
    assert _normalize_pid(f"P{_PID}") == _PID
    assert _normalize_pid("730-123-456") == _PID
    assert _normalize_pid(_APPROVAL_NUMBER) is None
    assert _normalize_pid("12345678") is None
    assert _normalize_pid("1234567890") is None
    assert _normalize_pid(1380307) is None
    assert _normalize_pid(730123456) == _PID


def test_dedupe_pids_drops_approval_number_beside_pid():
    assert _dedupe_pids([_PID, _APPROVAL_NUMBER, _PID]) == [_PID]
    assert _dedupe_pids([f"{_PID} {_APPROVAL_NUMBER}"]) == [_PID]
    assert _dedupe_pids([f"PID: {_PID} / {_APPROVAL_NUMBER}"]) == [_PID]
    assert _dedupe_pids([f"{_PID},{_APPROVAL_NUMBER}"]) == [_PID]


def test_ad_normalize_pid_rejects_approval_numbers():
    assert ad_normalize_pid(_PID) == _PID
    assert ad_normalize_pid(f"730-{_PID[3:6]}-{_PID[6:]}") == _PID
    assert ad_normalize_pid(_APPROVAL_NUMBER) == ""
    assert ad_normalize_pid("N/A") == ""


def test_preview_payload_drops_non_nine_digit_pids():
    payload = OutsideScholarshipService._build_preview_payload(
        filename="checks.pdf",
        aid_year="2027",
        aid_term="F",
        checks=[
            {
                "pid_list": [_PID, _APPROVAL_NUMBER, _PID, f"{_PID} {_APPROVAL_NUMBER}"],
                "amount": "1000.00",
                "check_number": "1001",
                "name": None,
                "provider": "Cabinetworks Group Michigan, LLC",
                "scholarship_name": None,
                "metadata": {"check_index": 1},
            }
        ],
        ad_by_pid={},
    )

    assert [entry["pid"] for entry in payload["checks"][0]["pids"]] == [_PID]


def test_excel_emits_one_row_when_extracted_has_pid_and_approval_number():
    """Approval numbers in extracted are filtered; reviewed keeps only the real PID."""
    loader = OutsideScholarshipsDataLoader(aid_year="2027", aid_term="F")
    rows = loader._expand_reviewed_rows(
        [
            {
                "reviewed": {
                    "pids": [{"pid": _PID}],
                    "amount": "1000.00",
                    "name": "",
                    "check_number": "1001",
                    "provider": "Cabinetworks Group Michigan, LLC",
                    "scholarship_name": "",
                },
                "extracted": {
                    "pid_list": [_PID, _APPROVAL_NUMBER],
                },
                "verified": True,
            }
        ]
    )

    assert len(rows) == 1
    assert rows[0][0] == _PID
    assert rows[0][1] == 1000.0
    assert rows[0][7] == "Cabinetworks Group Michigan, LLC"


def test_excel_keeps_multiple_nine_digit_pids_and_blank_when_none():
    loader = OutsideScholarshipsDataLoader(aid_year="2027", aid_term="F")
    rows = loader._expand_reviewed_rows(
        [
            {
                "reviewed": {"pids": [{"pid": _PID}, {"pid": _SECOND_PID}], "amount": "50.00"},
                "extracted": {},
            },
            {
                "reviewed": {"pids": [{"pid": ""}], "amount": "25.00"},
                "extracted": {"pid_list": [_APPROVAL_NUMBER]},
            },
        ]
    )

    assert [row[0] for row in rows] == [_PID, _SECOND_PID, ""]
    assert [row[1] for row in rows] == [50.0, 50.0, 25.0]


def test_build_reviewed_excel_rejects_non_nine_digit_reviewed_pid():
    loader = OutsideScholarshipsDataLoader(aid_year="2027", aid_term="F")
    with pytest.raises(ValueError, match=r"Check 1001: PID '73012345' is not 9 digits"):
        loader.build_reviewed_excel_bytes(
            {
                "checks": [
                    {
                        "check_index": 1,
                        "reviewed": {
                            "pids": [{"pid": "73012345"}],
                            "amount": "100.00",
                            "check_number": "1001",
                            "provider": "Provider",
                        },
                        "extracted": {"pids": [{"pid": _PID}]},
                        "verified": True,
                    }
                ]
            }
        )


def test_build_reviewed_excel_allows_blank_reviewed_pid():
    loader = OutsideScholarshipsDataLoader(aid_year="2027", aid_term="F")
    output = loader.build_reviewed_excel_bytes(
        {
            "checks": [
                {
                    "check_index": 1,
                    "reviewed": {
                        "pids": [{"pid": ""}],
                        "amount": "25.00",
                        "check_number": "1002",
                        "provider": "Provider",
                    },
                    "extracted": {},
                    "verified": True,
                }
            ]
        }
    )
    worksheet = load_workbook(output).active
    assert (worksheet.cell(row=7, column=1).value or "") == ""


def test_upload_with_approval_beside_pid_exports_one_excel_row(client, override_auth, app):
    """Issue #13: model (or page text) lists a shorter approval number next to the PID."""
    graph = FakeGraph()
    response, _ = upload(
        client,
        app,
        [FRONT, BACK],
        back_pids={1001: [_PID, _APPROVAL_NUMBER]},
        graph=graph,
    )

    assert response.status_code == 200, response.text
    preview = preview_payload(response)
    assert [entry["pid"] for entry in preview["checks"][0]["pids"]] == [_PID]
    assert graph.all_requested_pids() == [_PID]

    export_response = export_excel(client, app, preview, verified_indexes={1}, graph=graph)
    assert export_response.status_code == 200, export_response.text
    worksheet = load_workbook(BytesIO(export_response.content)).active
    excel_pids = [row[0] or "" for row in worksheet.iter_rows(min_row=REVIEW_FIRST_DATA_ROW, values_only=True)]
    assert excel_pids == [_PID]


def test_export_rejects_client_injected_approval_number(client, override_auth, app):
    """A non-9-digit reviewed PID must 400 — never silently drop a scholarship row."""
    from api.dependencies import get_azure_client, get_graph_user_directory
    from services.azure_services import GraphUserDirectory

    graph = FakeGraph()
    export_body = {
        "aid_year": "2027",
        "aid_term": "F",
        "checks": [
            {
                "check_index": 1,
                "extracted": {
                    "amount": "1000.00",
                    "check_number": "1001",
                    "name": "",
                    "provider": "Cabinetworks Group Michigan, LLC",
                    "scholarship_name": "",
                    "pids": [{"pid": _PID}],
                },
                "reviewed": {
                    "amount": "1000.00",
                    "check_number": "1001",
                    "name": "",
                    "provider": "Cabinetworks Group Michigan, LLC",
                    "scholarship_name": "",
                    "pids": [
                        {"pid": _PID, "active_directory": {"status": "not_found", "name": None}},
                        {
                            "pid": _APPROVAL_NUMBER,
                            "active_directory": {"status": "found", "name": "Fake Approval"},
                        },
                    ],
                },
                "verified": True,
            }
        ],
    }
    app.dependency_overrides[get_azure_client] = lambda: SimpleNamespace(llm=None)
    app.dependency_overrides[get_graph_user_directory] = lambda: GraphUserDirectory(transport=graph.transport)

    export_response = client.post("/api/banking/outside-scholarships/export", json=export_body)
    assert export_response.status_code == 400, export_response.text
    detail = export_response.json()["detail"]
    assert "1001" in detail
    assert _APPROVAL_NUMBER in detail


def test_export_rejects_eight_digit_reviewed_pid(client, override_auth, app):
    from api.dependencies import get_azure_client, get_graph_user_directory
    from services.azure_services import GraphUserDirectory

    graph = FakeGraph()
    export_body = {
        "aid_year": "2027",
        "aid_term": "F",
        "checks": [
            {
                "check_index": 1,
                "extracted": {
                    "amount": "100.00",
                    "check_number": "1001",
                    "provider": "Provider 1001",
                    "pids": [{"pid": _PID}],
                },
                "reviewed": {
                    "amount": "100.00",
                    "check_number": "1001",
                    "provider": "Provider 1001",
                    "pids": [{"pid": "73012345"}],
                },
                "verified": True,
            }
        ],
    }
    app.dependency_overrides[get_azure_client] = lambda: SimpleNamespace(llm=None)
    app.dependency_overrides[get_graph_user_directory] = lambda: GraphUserDirectory(transport=graph.transport)

    export_response = client.post("/api/banking/outside-scholarships/export", json=export_body)
    assert export_response.status_code == 400, export_response.text
    detail = export_response.json()["detail"]
    assert "1001" in detail
    assert "73012345" in detail
