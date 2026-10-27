"""Supporting pages stay with the check ahead of them and do not replace the check's money or PID."""

import logging
from io import BytesIO
from pathlib import Path

import fitz
import pytest
from openpyxl import load_workbook

from outside_scholarships_fakes import (
    BACK,
    FRONT,
    REVIEW_FIRST_DATA_ROW,
    SUPPORTING,
    export_excel,
    preview_payload,
    upload,
)
from prompts.outside_scholarship import (
    EXTRACT_OUTSIDE_SCHOLARSHIP_PROMPT,
    EXTRACT_SUPPORTING_ROSTER_PROMPT,
)
from services.banking.outside_scholarships.nodes import (
    PYMUPDF_LOCK,
    _force_first_page_front,
    _group_pages_by_check,
    _normalize_aid_term,
    _normalize_pid,
)


def test_grouping_checks_in_a_row_with_no_supporting_pages():
    assert _group_pages_by_check([FRONT, FRONT, BACK, FRONT]) == [
        {"front": 0, "back": None, "supporting": []},
        {"front": 1, "back": 2, "supporting": []},
        {"front": 3, "back": None, "supporting": []},
    ]


def test_grouping_check_with_several_supporting_pages():
    assert _group_pages_by_check([FRONT, SUPPORTING, SUPPORTING, FRONT, SUPPORTING]) == [
        {"front": 0, "back": None, "supporting": [1, 2]},
        {"front": 3, "back": None, "supporting": [4]},
    ]


def test_grouping_front_back_then_supporting():
    assert _group_pages_by_check([FRONT, BACK, SUPPORTING, SUPPORTING]) == [
        {"front": 0, "back": 1, "supporting": [2, 3]},
    ]


def test_grouping_rejects_a_supporting_page_before_any_check():
    with pytest.raises(ValueError, match="no check before it"):
        _group_pages_by_check([SUPPORTING, FRONT])


def test_force_first_page_front_overrides_a_misclassified_page_one():
    assert _force_first_page_front([SUPPORTING, FRONT, SUPPORTING]) == [FRONT, FRONT, SUPPORTING]
    assert _force_first_page_front([FRONT, BACK]) == [FRONT, BACK]


def test_misclassified_page_one_still_uploads(client, override_auth, app):
    """Classifier can call page 1 supporting; pairing forces front so the PDF is not rejected."""
    response, _ = upload(client, app, [SUPPORTING, FRONT, BACK], back_pids={1002: "730000002"})
    assert response.status_code == 200, response.text
    checks = preview_payload(response)["checks"]
    assert [check["front_page"] for check in checks] == [1, 2]
    assert checks[0]["supporting_pages"] == []
    assert checks[1]["back_page"] == 3


def test_grouping_keeps_a_back_that_was_scanned_after_a_supporting_page():
    assert _group_pages_by_check([FRONT, SUPPORTING, BACK, SUPPORTING]) == [
        {"front": 0, "back": 2, "supporting": [1, 3]},
    ]


def test_a_back_scanned_after_a_stub_is_still_the_check_back(client, override_auth, app):
    response, llm = upload(
        client,
        app,
        [FRONT, SUPPORTING, BACK],
        back_pids={1001: "730000001"},
    )
    assert response.status_code == 200, response.text
    check = preview_payload(response)["checks"][0]
    assert check["front_page"] == 1
    assert check["back_page"] == 3
    assert check["supporting_pages"] == [2]
    assert [entry["pid"] for entry in check["pids"]] == ["730000001"]
    assert llm.image_counts_by_check()[1001] == 3


def test_grouping_rejects_a_second_back():
    with pytest.raises(ValueError, match="only one back"):
        _group_pages_by_check([FRONT, BACK, BACK])


def test_normalize_aid_term_drops_wording_that_is_not_a_code():
    assert _normalize_aid_term("F/S") == "F/S"
    assert _normalize_aid_term(" ss2 ") == "SS2"
    assert _normalize_aid_term("Fall 2026 and Spring 2027") is None
    assert _normalize_aid_term(None) is None


def test_normalize_pid_accepts_json_whole_numbers():
    assert _normalize_pid(730123456) == "730123456"
    assert _normalize_pid(730123456.0) == "730123456"
    assert _normalize_pid("730123456") == "730123456"
    assert _normalize_pid("730-123-456") == "730123456"
    assert _normalize_pid("000123456") is None
    assert _normalize_pid("000000001") is None
    assert _normalize_pid(True) is None
    assert _normalize_pid(False) is None
    assert _normalize_pid(73012345) is None


def test_prompt_ignores_an_unmarked_option_list():
    assert "if not indicated, we will split evenly" in EXTRACT_OUTSIDE_SCHOLARSHIP_PROMPT
    assert "printed list of options is not a selection" in EXTRACT_OUTSIDE_SCHOLARSHIP_PROMPT
    assert "do not add that stray PID to pid_list" in EXTRACT_OUTSIDE_SCHOLARSHIP_PROMPT
    assert "supporting roster" in EXTRACT_OUTSIDE_SCHOLARSHIP_PROMPT
    assert "supporting_has_amount_and_pid" in EXTRACT_OUTSIDE_SCHOLARSHIP_PROMPT
    assert EXTRACT_SUPPORTING_ROSTER_PROMPT.splitlines()[0] == "Extract every student row"
    assert "exactly nine digits you can actually see" in EXTRACT_SUPPORTING_ROSTER_PROMPT
    assert "never padded with leading zeros" in EXTRACT_SUPPORTING_ROSTER_PROMPT
    assert "never a row index, dollar amount, or other column" in EXTRACT_SUPPORTING_ROSTER_PROMPT
    assert "omit the whole row" in EXTRACT_SUPPORTING_ROSTER_PROMPT
    assert "Do not guess or invent PIDs" in EXTRACT_SUPPORTING_ROSTER_PROMPT
    assert "spreadsheet on its side" in EXTRACT_SUPPORTING_ROSTER_PROMPT
    assert "one object per" in EXTRACT_SUPPORTING_ROSTER_PROMPT
    assert "amount" in EXTRACT_SUPPORTING_ROSTER_PROMPT
    assert "JSON string with cents" in EXTRACT_SUPPORTING_ROSTER_PROMPT
    assert '"amount": null' in EXTRACT_SUPPORTING_ROSTER_PROMPT
    assert "must not change `amount` or `pid_list`" not in EXTRACT_OUTSIDE_SCHOLARSHIP_PROMPT


def test_one_check_extraction_failure_does_not_fail_the_upload(client, override_auth, app):
    import json

    def extraction_reply(check_number, fields):
        if check_number == 1001:
            raise RuntimeError("simulated vision failure")
        return json.dumps(fields)

    response, _ = upload(
        client,
        app,
        [FRONT, BACK, FRONT, BACK],
        back_pids={1001: "730000001", 1002: "730000002"},
        extraction_reply=extraction_reply,
    )
    assert response.status_code == 200, response.text
    first, second = preview_payload(response)["checks"]
    assert first["extraction_failed"] is True
    assert first["amount"] is None
    assert second["extraction_failed"] is False
    assert second["amount"] == "200.00"
    assert [entry["pid"] for entry in second["pids"]] == ["730000002"]


def test_supporting_pages_fill_name_scholarship_and_term_without_replacing_amount_or_pid(
    client, override_auth, app
):
    response, llm = upload(
        client,
        app,
        [FRONT, BACK, SUPPORTING, SUPPORTING, FRONT],
        back_pids={1001: "730000001"},
        supporting_by_page={
            3: {
                "scholarship_name": "Wells Fargo Employees Scholarship",
                "student_name": "Ella Kate DeLapp",
                "aid_term": "F/S",
                "supporting_amount": "9999.00",
                "supporting_pid": "730999999",
            },
            4: {
                "aid_term": None,
                "supporting_amount": "10000.00",
            },
        },
        aid_term="F",
    )
    assert response.status_code == 200, response.text
    preview = preview_payload(response)
    first, second = preview["checks"]
    assert first["amount"] == "100.00"
    assert [entry["pid"] for entry in first["pids"]] == ["730000001"]
    assert first["name"] == "Ella Kate DeLapp"
    assert first["scholarship_name"] == "Wells Fargo Employees Scholarship"
    assert first["aid_term"] == "F/S"
    assert first["supporting_pages"] == [3, 4]
    assert len(first["supporting_images"]) == 2
    assert all(image.startswith("data:image/jpeg;base64,") for image in first["supporting_images"])
    assert first["supporting_has_amount_and_pid"] is True
    assert first["extraction_failed"] is False
    assert llm.image_counts_by_check()[1001] == 4
    assert second["aid_term"] is None
    assert second["supporting_images"] == []
    assert second["amount"] == "200.00"
    assert second["supporting_has_amount_and_pid"] is False

    export_response = export_excel(client, app, preview)
    assert export_response.status_code == 200, export_response.text
    worksheet = load_workbook(BytesIO(export_response.content)).active
    rows = list(worksheet.iter_rows(min_row=REVIEW_FIRST_DATA_ROW, values_only=True))
    assert rows[0][6] == "F/S"
    assert rows[1][6] == "F"


def _excel_amounts(response):
    worksheet = load_workbook(BytesIO(response.content)).active
    return [
        row[1]
        for row in worksheet.iter_rows(min_row=REVIEW_FIRST_DATA_ROW, values_only=True)
    ]


def test_roster_supporting_pages_fill_pid_list_and_excel_rows(client, override_auth, app):
    response, llm = upload(
        client,
        app,
        [FRONT, SUPPORTING],
        back_pids={},
        supporting_by_page={
            2: {
                "supporting_roster": [
                    {
                        "pid": "730000001",
                        "amount": "30.00",
                        "name": "Test Student One",
                        "scholarship_name": "Roster Award One",
                    },
                    {
                        "pid": "730000002",
                        "amount": "30.00",
                        "name": "Test Student Two",
                        "scholarship_name": "Roster Award Two",
                    },
                    {
                        "pid": "730000003",
                        "amount": "40.00",
                        "name": "Test Student Three",
                        "scholarship_name": "Roster Award Three",
                    },
                    {
                        "pid": "000123456",
                        "amount": "625.00",
                        "name": "Padded Row Number",
                        "scholarship_name": "Not A PID",
                    },
                ]
            },
        },
    )
    assert response.status_code == 200, response.text
    assert llm.roster_calls == 3
    with PYMUPDF_LOCK, fitz.open() as blank:
        page_rect = blank.new_page().rect
    scale = 250 / 72
    rendered = (page_rect * fitz.Matrix(scale, scale)).irect
    top_height = round(rendered.height * (1 / 3 + 0.01))
    middle_start = round(rendered.height * (1 / 3 - 0.01))
    middle_end = round(rendered.height * (2 / 3 + 0.01))
    bottom_start = round(rendered.height * (2 / 3 - 0.01))
    assert llm.roster_image_sizes == [
        (rendered.width, top_height),
        (rendered.width, middle_end - middle_start),
        (rendered.width, rendered.height - bottom_start),
    ]

    preview = preview_payload(response)
    check = preview["checks"][0]
    assert check["check_number"] == "1001"
    assert [(entry["pid"], entry["amount"]) for entry in check["pids"]] == [
        ("730000001", "30.00"),
        ("730000002", "30.00"),
        ("730000003", "40.00"),
    ]
    assert [(entry["name"], entry["scholarship_name"]) for entry in check["pids"]] == [
        ("Test Student One", "Roster Award One"),
        ("Test Student Two", "Roster Award Two"),
        ("Test Student Three", "Roster Award Three"),
    ]

    export_response = export_excel(client, app, preview)
    assert export_response.status_code == 200, export_response.text
    artifact_path = Path(__file__).parent / "artifacts" / "outside-scholarships-roster-pids.xlsx"
    artifact_path.parent.mkdir(exist_ok=True)
    artifact_path.write_bytes(export_response.content)

    worksheet = load_workbook(BytesIO(export_response.content)).active
    rows = list(worksheet.iter_rows(min_row=REVIEW_FIRST_DATA_ROW, values_only=True))
    assert [(row[0], row[1], row[2], row[3], row[8]) for row in rows] == [
        ("730000001", 30, "1001", "Test Student One", "Roster Award One"),
        ("730000002", 30, "1001", "Test Student Two", "Roster Award Two"),
        ("730000003", 40, "1001", "Test Student Three", "Roster Award Three"),
    ]


def test_first_extract_one_pid_equal_to_check_still_reads_roster(client, override_auth, app):
    import json

    def extraction_reply(check_number, fields):
        fields = dict(fields)
        fields["pid_list"] = ["730000001"]
        fields["pid_amounts"] = [{"pid": "730000001", "amount": "100.00"}]
        return json.dumps(fields)

    response, llm = upload(
        client,
        app,
        [FRONT, SUPPORTING],
        back_pids={},
        supporting_by_page={
            2: {
                "supporting_roster": [
                    {
                        "pid": "730000001",
                        "amount": "30.00",
                        "name": "Test Student One",
                        "scholarship_name": "Roster Award One",
                    },
                    {
                        "pid": "730000002",
                        "amount": "30.00",
                        "name": "Test Student Two",
                        "scholarship_name": "Roster Award Two",
                    },
                    {
                        "pid": "730000003",
                        "amount": "40.00",
                        "name": "Test Student Three",
                        "scholarship_name": "Roster Award Three",
                    },
                ]
            },
        },
        extraction_reply=extraction_reply,
    )
    assert response.status_code == 200, response.text
    assert llm.roster_calls == 3

    preview = preview_payload(response)
    check = preview["checks"][0]
    assert check["check_number"] == "1001"
    assert [(entry["pid"], entry["amount"]) for entry in check["pids"]] == [
        ("730000001", "30.00"),
        ("730000002", "30.00"),
        ("730000003", "40.00"),
    ]

    export_response = export_excel(client, app, preview)
    assert export_response.status_code == 200, export_response.text
    assert _excel_amounts(export_response) == [30, 30, 40]


def test_complete_supporting_pid_amount_split_reaches_preview_and_excel(client, override_auth, app):
    response, _ = upload(
        client,
        app,
        [FRONT, BACK, SUPPORTING, SUPPORTING, SUPPORTING, SUPPORTING],
        back_pids={1001: ["730000001", "730000002", "730000003"]},
        supporting_by_page={
            3: {"supporting_pid": "730000001", "supporting_amount": "30.00"},
            4: {"supporting_pid": "730000002", "supporting_amount": "30.00"},
            5: {"supporting_pid": "730000003", "supporting_amount": "40.00"},
            6: {"supporting_pid": "730999999", "supporting_amount": "999.00"},
        },
    )
    assert response.status_code == 200, response.text

    preview = preview_payload(response)
    pids = preview["checks"][0]["pids"]
    assert [(entry["pid"], entry["amount"]) for entry in pids] == [
        ("730000001", "30.00"),
        ("730000002", "30.00"),
        ("730000003", "40.00"),
    ]

    export_response = export_excel(client, app, preview)
    assert export_response.status_code == 200, export_response.text
    assert _excel_amounts(export_response) == [30, 30, 40]


def test_roster_stray_pid_on_a_letter_is_still_ignored(client, override_auth, app):
    response, llm = upload(
        client,
        app,
        [FRONT, BACK, SUPPORTING, SUPPORTING, SUPPORTING, SUPPORTING],
        back_pids={1001: ["730000001", "730000002", "730000003"]},
        supporting_by_page={
            3: {"supporting_pid": "730000001", "supporting_amount": "30.00"},
            4: {"supporting_pid": "730000002", "supporting_amount": "30.00"},
            5: {"supporting_pid": "730000003", "supporting_amount": "40.00"},
            6: {"supporting_pid": "730999999", "supporting_amount": "999.00"},
        },
    )
    assert response.status_code == 200, response.text
    assert llm.roster_calls == 0

    preview = preview_payload(response)
    assert [entry["pid"] for entry in preview["checks"][0]["pids"]] == [
        "730000001",
        "730000002",
        "730000003",
    ]


def test_incomplete_roster_exports_full_check_amount(client, override_auth, app, caplog):
    caplog.set_level(logging.INFO, logger="services.banking.outside_scholarships.nodes")
    response, llm = upload(
        client,
        app,
        [FRONT, SUPPORTING],
        back_pids={},
        supporting_by_page={
            2: {
                "supporting_roster": [
                    {"pid": "730000001", "amount": "30.00"},
                    {"pid": "730000002", "amount": "30.00"},
                    {"pid": "730000003", "amount": "30.00"},
                ]
            },
        },
    )
    assert response.status_code == 200, response.text
    assert llm.roster_calls == 3
    assert (
        "outside_scholarships.extract_roster_pages: crop_count=3 retry_count=0 rotate=0 "
        "kept_pid_count=3 kept_amount_count=3"
    ) in caplog.text

    preview = preview_payload(response)
    assert [(entry["pid"], entry.get("amount")) for entry in preview["checks"][0]["pids"]] == [
        ("730000001", "30.00"),
        ("730000002", "30.00"),
        ("730000003", "30.00"),
    ]

    export_response = export_excel(client, app, preview)
    assert export_response.status_code == 200, export_response.text
    assert _excel_amounts(export_response) == [100, 100, 100]


def test_roster_supporting_pages_are_read_one_page_at_a_time(client, override_auth, app):
    response, llm = upload(
        client,
        app,
        [FRONT, SUPPORTING, SUPPORTING],
        back_pids={},
        supporting_by_page={
            2: {
                "supporting_roster": [
                    {"pid": "730000001", "amount": "30.00"},
                    {"pid": "730000002", "amount": "30.00"},
                ]
            },
            3: {
                "supporting_roster": [
                    {"pid": "730000003", "amount": "40.00"},
                ]
            },
        },
    )
    assert response.status_code == 200, response.text
    assert llm.roster_calls == 6

    check = preview_payload(response)["checks"][0]
    assert check["check_number"] == "1001"
    assert [(entry["pid"], entry["amount"]) for entry in check["pids"]] == [
        ("730000001", "30.00"),
        ("730000002", "30.00"),
        ("730000003", "40.00"),
    ]


def test_roster_retry_empty_page_when_another_page_had_rows(client, override_auth, app):
    import json

    def roster_reply(page, rows, call_count):
        if page.page_number == 3 and call_count == 1:
            return json.dumps({"rows": []})
        return json.dumps({"rows": rows})

    response, llm = upload(
        client,
        app,
        [FRONT, SUPPORTING, SUPPORTING],
        back_pids={},
        supporting_by_page={
            2: {
                "supporting_roster": [
                    {"pid": "730000001", "amount": "60.00"},
                    {"pid": "730000002", "amount": "40.00"},
                ]
            },
            3: {
                "supporting_roster": [
                    {"pid": "730000003", "amount": "20.00"},
                    {"pid": "730000004", "amount": "80.00"},
                ]
            },
        },
        roster_reply=roster_reply,
    )
    assert response.status_code == 200, response.text
    assert llm.roster_calls == 7

    check = preview_payload(response)["checks"][0]
    assert [(entry["pid"], entry["amount"]) for entry in check["pids"]] == [
        ("730000001", "60.00"),
        ("730000002", "40.00"),
        ("730000003", "20.00"),
        ("730000004", "80.00"),
    ]


def test_roster_followup_ignores_supporting_letter_pages(client, override_auth, app):
    response, llm = upload(
        client,
        app,
        [FRONT, SUPPORTING, SUPPORTING],
        back_pids={},
        supporting_by_page={
            2: {
                "supporting_roster": [
                    {"pid": "730000001", "amount": "60.00"},
                    {"pid": "730000002", "amount": "40.00"},
                ]
            },
            3: {
                "supporting_pid": "730999999",
                "supporting_amount": "999.00",
            },
        },
    )
    assert response.status_code == 200, response.text
    assert llm.roster_calls == 9

    check = preview_payload(response)["checks"][0]
    assert [(entry["pid"], entry["amount"]) for entry in check["pids"]] == [
        ("730000001", "60.00"),
        ("730000002", "40.00"),
    ]


def test_roster_retries_rotated_page_when_first_pass_has_no_amounts(client, override_auth, app):
    import json

    def roster_reply(page, rows, call_count):
        if call_count <= 3:
            return json.dumps({"rows": [{**row, "amount": None} for row in rows]})
        return json.dumps({"rows": rows})

    response, llm = upload(
        client,
        app,
        [FRONT, SUPPORTING],
        back_pids={},
        supporting_by_page={
            2: {
                "supporting_roster": [
                    {"pid": "730000001", "amount": "30.00"},
                    {"pid": "730000002", "amount": "30.00"},
                    {"pid": "730000003", "amount": "40.00"},
                ]
            },
        },
        roster_reply=roster_reply,
    )
    assert response.status_code == 200, response.text
    assert llm.roster_calls == 7

    check = preview_payload(response)["checks"][0]
    assert [(entry["pid"], entry["amount"]) for entry in check["pids"]] == [
        ("730000001", "30.00"),
        ("730000002", "30.00"),
        ("730000003", "40.00"),
    ]

    export_response = export_excel(client, app, preview_payload(response))
    assert export_response.status_code == 200, export_response.text
    assert _excel_amounts(export_response) == [30, 30, 40]


def test_partial_first_extract_still_runs_roster_followup(client, override_auth, app):
    import json

    def extraction_reply(check_number, fields):
        fields = dict(fields)
        fields["pid_list"] = ["730000001", "730000002"]
        fields["pid_amounts"] = [
            {"pid": "730000001", "amount": "30.00"},
            {"pid": "730000002", "amount": "30.00"},
        ]
        return json.dumps(fields)

    response, llm = upload(
        client,
        app,
        [FRONT, SUPPORTING],
        back_pids={},
        supporting_by_page={
            2: {
                "supporting_roster": [
                    {"pid": "730000001", "amount": "30.00"},
                    {"pid": "730000002", "amount": "30.00"},
                    {"pid": "730000003", "amount": "40.00"},
                ]
            },
        },
        extraction_reply=extraction_reply,
    )
    assert response.status_code == 200, response.text
    assert llm.roster_calls == 3

    check = preview_payload(response)["checks"][0]
    assert [(entry["pid"], entry["amount"]) for entry in check["pids"]] == [
        ("730000001", "30.00"),
        ("730000002", "30.00"),
        ("730000003", "40.00"),
    ]


def test_incomplete_supporting_pid_amount_split_exports_full_check_amount(client, override_auth, app):
    response, _ = upload(
        client,
        app,
        [FRONT, BACK, SUPPORTING, SUPPORTING],
        back_pids={1001: ["730000001", "730000002", "730000003"]},
        supporting_by_page={
            3: {"supporting_pid": "730000001", "supporting_amount": "30.00"},
            4: {"supporting_pid": "730000002", "supporting_amount": "30.00"},
        },
    )
    assert response.status_code == 200, response.text

    preview = preview_payload(response)
    assert [entry.get("amount") for entry in preview["checks"][0]["pids"]] == [
        "30.00",
        "30.00",
        None,
    ]

    export_response = export_excel(client, app, preview)
    assert export_response.status_code == 200, export_response.text
    assert _excel_amounts(export_response) == [100, 100, 100]


def test_mismatched_supporting_pid_amount_split_exports_full_check_amount(client, override_auth, app):
    response, _ = upload(
        client,
        app,
        [FRONT, BACK, SUPPORTING, SUPPORTING, SUPPORTING],
        back_pids={1001: ["730000001", "730000002", "730000003"]},
        supporting_by_page={
            3: {"supporting_pid": "730000001", "supporting_amount": "30.00"},
            4: {"supporting_pid": "730000002", "supporting_amount": "30.00"},
            5: {"supporting_pid": "730000003", "supporting_amount": "30.00"},
        },
    )
    assert response.status_code == 200, response.text

    preview = preview_payload(response)
    export_response = export_excel(client, app, preview)
    assert export_response.status_code == 200, export_response.text
    assert _excel_amounts(export_response) == [100, 100, 100]
