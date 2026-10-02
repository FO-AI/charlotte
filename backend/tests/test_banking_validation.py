from io import BytesIO
from unittest.mock import MagicMock

from openpyxl import load_workbook

from outside_scholarships_fakes import FRONT, REVIEW_FIRST_DATA_ROW, export_excel, preview_payload, upload
from services.data_loaders.outside_scholarships_json_to_excel import ALLOWED_AID_TERMS


def test_outside_scholarships_rejects_non_pdf(client, override_auth, app):
    """Validation returns 400; Depends still injects a client, but it is unused."""
    from api.dependencies import get_azure_client

    azure = MagicMock()
    app.dependency_overrides[get_azure_client] = lambda: azure

    response = client.post(
        "/api/banking/outside-scholarships",
        files={"files": ("notes.txt", b"not a pdf", "text/plain")},
        data={"aid_term": "F"},
    )
    assert response.status_code == 400
    assert "PDF" in response.json()["detail"]
    # FastAPI resolves Depends before the handler; the mock is injected but never used.
    assert azure.mock_calls == []


def test_outside_scholarships_rejects_invalid_aid_term(client, override_auth, app):
    from api.dependencies import get_azure_client

    app.dependency_overrides[get_azure_client] = lambda: MagicMock()
    response = client.post(
        "/api/banking/outside-scholarships",
        files={"files": ("checks.pdf", b"%PDF-1.1", "application/pdf")},
        data={"aid_term": "X"},
    )
    assert response.status_code == 400
    detail = response.json()["detail"]
    assert "Aid term" in detail
    for term in ALLOWED_AID_TERMS:
        assert term in detail


def test_outside_scholarships_accepts_ss1_aid_term_through_export(client, override_auth, app):
    response, _ = upload(client, app, [FRONT], aid_term="ss1")
    assert response.status_code == 200, response.text
    preview = preview_payload(response)
    assert preview["aid_term"] == "SS1"

    export_response = export_excel(client, app, preview, verified_indexes={1})
    assert export_response.status_code == 200, export_response.text
    worksheet = load_workbook(BytesIO(export_response.content)).active
    rows = list(worksheet.iter_rows(min_row=REVIEW_FIRST_DATA_ROW, values_only=True))
    assert rows[0][6] == "SS1"
