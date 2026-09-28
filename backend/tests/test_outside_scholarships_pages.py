"""E2E: outside-scholarship uploads where some checks have no scanned back.

No real Azure access. Document Intelligence and the LLM are fakes that read the text
PyMuPDF writes onto each test page, so every assertion reflects which pages actually
reached each check: fronts carry the check number and amount, backs carry the PID.
"""

import json

import pytest
from openpyxl import load_workbook

from outside_scholarships_fakes import BACK, FIRST_DATA_ROW, FRONT, upload

def _excel_rows(response, tmp_path):
    """Save the returned workbook for inspection and return (PID, amount) per data row."""
    workbook_path = tmp_path / "outside_scholarships.xlsx"
    workbook_path.write_bytes(response.content)
    worksheet = load_workbook(workbook_path).active
    return [
        (pid or "", amount)
        for pid, amount, *_ in worksheet.iter_rows(min_row=FIRST_DATA_ROW, values_only=True)
    ]


@pytest.mark.parametrize(
    ("sides", "expected_rows", "expected_image_counts"),
    [
        pytest.param(
            [FRONT],
            [("", 100.0)],
            {"1001": 1},
            id="single-front",
        ),
        pytest.param(
            [FRONT, BACK, FRONT],
            [("P0001001", 100.0), ("", 200.0)],
            {"1001": 2, "1002": 1},
            id="odd-page-count",
        ),
        pytest.param(
            [FRONT, BACK, FRONT, FRONT, BACK],
            [("P0001001", 100.0), ("", 200.0), ("P0001003", 300.0)],
            {"1001": 2, "1002": 1, "1003": 2},
            id="missing-back-in-middle",
        ),
        pytest.param(
            [FRONT, FRONT, FRONT, FRONT],
            [("", 100.0), ("", 200.0), ("", 300.0), ("", 400.0)],
            {"1001": 1, "1002": 1, "1003": 1, "1004": 1},
            id="all-fronts-even-count",
        ),
        pytest.param(
            [FRONT, BACK, FRONT, BACK],
            [("P0001001", 100.0), ("P0001002", 200.0)],
            {"1001": 2, "1002": 2},
            id="every-check-has-back",
        ),
    ],
)
def test_pages_group_into_checks_by_detected_side(
    client, override_auth, app, tmp_path, sides, expected_rows, expected_image_counts
):
    response, llm, _ = upload(client, app, sides)

    assert response.status_code == 200, response.text
    assert _excel_rows(response, tmp_path) == expected_rows
    assert llm.verify_image_counts == expected_image_counts


@pytest.mark.parametrize(
    ("sides", "orphan_page"),
    [
        pytest.param([BACK, FRONT], 1, id="leading-back"),
        pytest.param([FRONT, BACK, BACK], 3, id="two-backs-in-a-row"),
    ],
)
def test_back_without_front_is_rejected_with_page_number(client, override_auth, app, sides, orphan_page):
    response, _, document_intelligence = upload(client, app, sides)

    assert response.status_code == 400
    assert f"Page {orphan_page} " in response.json()["detail"]
    assert document_intelligence.calls == 0


@pytest.mark.parametrize(
    "classify_reply",
    [
        pytest.param(lambda pages: "not json", id="not-json"),
        pytest.param(
            lambda pages: json.dumps({"pages": [{"page": page, "side": FRONT} for page in pages[:-1]]}),
            id="page-missing",
        ),
        pytest.param(
            lambda pages: json.dumps({"pages": [{"page": page, "side": "sideways"} for page in pages]}),
            id="unknown-side",
        ),
    ],
)
def test_invalid_classifier_reply_fails_before_extraction(client, override_auth, app, classify_reply):
    response, _, document_intelligence = upload(client, app, [FRONT, BACK], classify_reply)

    assert response.status_code == 500
    assert document_intelligence.calls == 0


def test_classifier_formatting_variations_are_accepted(client, override_auth, app, tmp_path):
    sides = [FRONT, BACK, FRONT]

    def loosely_formatted_reply(pages):
        return json.dumps(
            {"pages": [{"page": str(page), "side": f" {sides[page - 1].title()} "} for page in pages]}
        )

    response, _, _ = upload(client, app, sides, loosely_formatted_reply)

    assert response.status_code == 200, response.text
    assert _excel_rows(response, tmp_path) == [("P0001001", 100.0), ("", 200.0)]


def test_pages_beyond_one_batch_keep_pdf_page_numbers(client, override_auth, app, tmp_path):
    sides = [FRONT, BACK] * 6

    response, llm, _ = upload(client, app, sides)

    assert response.status_code == 200, response.text
    # Batches are sent concurrently, so compare them independent of arrival order.
    assert sorted(llm.classified_pages) == [list(range(1, 11)), [11, 12]]
    assert [pid for pid, _ in _excel_rows(response, tmp_path)] == [f"P{1000 + i:07d}" for i in range(1, 7)]
